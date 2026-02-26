import os
import re
import random
from datetime import datetime
from flask import Flask, request, jsonify, render_template
from flask_cors import CORS

# File parsing
import pdfplumber
import docx
import openpyxl
from PIL import Image
import pytesseract

UPLOAD_FOLDER = "uploads"
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt", "xlsx", "csv", "png", "jpg", "jpeg"}

app = Flask(__name__)
CORS(app)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# ----------------------------
# Helper Functions
# ----------------------------
def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text(file_path):
    ext = file_path.rsplit(".",1)[1].lower()
    text = ""

    try:
        if ext == "txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        elif ext == "pdf":
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        elif ext == "docx":
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif ext in ["xlsx", "csv"]:
            wb = openpyxl.load_workbook(file_path) if ext=="xlsx" else None
            if wb:
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        text += " ".join([str(cell) for cell in row if cell]) + "\n"
        elif ext in ["png","jpg","jpeg"]:
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img)
    except Exception as e:
        text += f"\n[Error extracting text: {str(e)}]"

    return text

def detect_sensitive_data(text):
    issues = []

    # Emails
    emails = re.findall(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text)
    if emails: issues.append(f"Emails found: {len(emails)}")

    # Phone numbers
    phones = re.findall(r"\+?\d[\d\s-]{7,}\d", text)
    if phones: issues.append(f"Phone numbers found: {len(phones)}")

    # Credit cards
    cards = re.findall(r"\b(?:\d[ -]*?){13,16}\b", text)
    if cards: issues.append(f"Credit cards found: {len(cards)}")

    # IDs
    ids = re.findall(r"\b\d{6,12}\b", text)
    if ids: issues.append(f"IDs found: {len(ids)}")

    # API Keys
    apis = re.findall(r"(?i)api[_-]?key[=:]\w+", text)
    if apis: issues.append(f"API keys found: {len(apis)}")

    return issues

def generate_risk_score(issues):
    base = 100
    deduction = len(issues)*15
    score = max(0, base-deduction)
    level = "Low" if score>70 else "Medium" if score>40 else "High"
    return score, level

def generate_recommendations(issues):
    recs=[]
    if not issues:
        recs.append("No sensitive data detected. Maintain security policies.")
    else:
        if any("Email" in i for i in issues):
            recs.append("Redact or encrypt emails before sharing files.")
        if any("Phone" in i for i in issues):
            recs.append("Remove or encrypt phone numbers.")
        if any("Credit" in i for i in issues):
            recs.append("Encrypt files containing financial data.")
        if any("ID" in i for i in issues):
            recs.append("Mask or remove IDs from shared documents.")
        if any("API" in i for i in issues):
            recs.append("Rotate API keys and store securely.")
    return recs

# ----------------------------
# Routes
# ----------------------------
@app.route("/")
def dashboard():
    return render_template("dashboard.html")

@app.route("/api/scan-document", methods=["POST"])
def scan_document():
    if "document" not in request.files:
        return jsonify({"error":"No document uploaded"}),400
    file = request.files["document"]
    if file.filename=="":
        return jsonify({"error":"No file selected"}),400
    if not allowed_file(file.filename):
        return jsonify({"error":"File type not allowed"}),400

    filename = f"{int(datetime.now().timestamp())}_{file.filename}"
    file_path = os.path.join(app.config["UPLOAD_FOLDER"],filename)
    file.save(file_path)

    text = extract_text(file_path)
    issues = detect_sensitive_data(text)
    risk_score, risk_level = generate_risk_score(issues)
    recommendations = generate_recommendations(issues)

    response={
        "filename":file.filename,
        "scannedAt":datetime.now().isoformat(),
        "issuesFound":issues,
        "complianceScore":risk_score,
        "riskLevel":risk_level,
        "recommendations":recommendations,
        "status":"Completed"
    }
    return jsonify(response)

@app.route("/health")
def health_check():
    return jsonify({"message":"PrivGuard AI Backend Running"})

# ----------------------------
# Run server
# ----------------------------
if __name__=="__main__":
    app.run(debug=True,port=5000)